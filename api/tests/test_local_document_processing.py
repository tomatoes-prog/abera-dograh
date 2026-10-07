"""Real local readers and OCR, exercised without any network or provider key."""

import asyncio
import json
import zipfile
from pathlib import Path
from unittest.mock import AsyncMock

import pytest
from arq import Retry
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from api.services.knowledge_base import processing


def _text_pdf(path: Path):
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
            NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    content = DecodedStreamObject()
    content.set_data(
        b"BT /F1 18 Tf 40 700 Td (Colombia: atenci\xf3n al cliente, precio 25000 pesos.) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.write(path)


@pytest.mark.parametrize("extension", [".txt", ".md", ".json"])
@pytest.mark.parametrize("mode", ["chunked", "full_document"])
def test_text_formats_preserve_spanish_and_retrieval_modes(tmp_path, extension, mode):
    text = "Bogotá: atención, envío y devolución. 😊\n" * 50
    path = tmp_path / ("catálogo" + extension)
    path.write_text(
        json.dumps({"productos": text}, ensure_ascii=False)
        if extension == ".json"
        else text,
        encoding="utf-8",
    )
    result = processing.process_document(path, path.name, mode, 32)
    assert "Bogotá" in result["full_text"]
    assert "devolución" in result["full_text"]
    assert result["docling_metadata"]["processor"] == "abera-local"
    if mode == "full_document":
        assert result["chunks"] == []
    else:
        assert (
            "".join(chunk["chunk_text"] for chunk in result["chunks"])
            == result["full_text"]
        )
        assert all(0 < chunk["token_count"] <= 32 for chunk in result["chunks"])
        assert "�" not in "".join(chunk["chunk_text"] for chunk in result["chunks"])


def test_windows_unicode_export(tmp_path):
    path = tmp_path / "exportación.txt"
    path.write_text("Devolución en Bogotá", encoding="utf-16")
    text, _ = processing.extract_document(path, path.name)
    assert text == "Devolución en Bogotá"


def test_docx_keeps_headings_tables_and_paragraphs(tmp_path):
    path = tmp_path / "catálogo.docx"
    document = Document()
    document.add_heading("Productos colombianos", 1)
    document.add_paragraph("Envíos a Bogotá y Medellín.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Producto"
    table.cell(0, 1).text = "Precio COP"
    table.cell(1, 0).text = "Café"
    table.cell(1, 1).text = "25.000"
    document.save(path)
    text, _ = processing.extract_document(path, path.name)
    assert "# Productos colombianos" in text
    assert "Envíos a Bogotá y Medellín." in text
    assert "Café | 25.000" in text


def test_pdf_text_and_ocr(tmp_path):
    text_path = tmp_path / "digital.pdf"
    _text_pdf(text_path)
    text, metadata = processing.extract_document(text_path, text_path.name)
    assert "atención al cliente" in text
    assert "25000" in text
    assert metadata["ocr_pages"] == []

    image = Image.new("RGB", (1200, 400), "white")
    draw = ImageDraw.Draw(image)
    draw.text(
        (60, 60),
        "Colombia - pedidos y soporte",
        font=ImageFont.load_default(size=40),
        fill="black",
    )
    draw.text(
        (60, 140),
        "Precio: 25000 pesos",
        font=ImageFont.load_default(size=40),
        fill="black",
    )
    scan_path = tmp_path / "escaneado.pdf"
    image.save(scan_path, "PDF")
    text, metadata = processing.extract_document(scan_path, scan_path.name)
    assert "Colombia" in text and "25000" in text
    assert metadata["ocr_pages"] == [1]


def test_rejects_unreadable_and_oversized_inputs(tmp_path):
    path = tmp_path / "bad.json"
    path.write_text("{broken", encoding="utf-8")
    with pytest.raises(ValueError):
        processing.extract_document(path, path.name)
    with pytest.raises(ValueError, match="Formato"):
        processing.extract_document(path, "script.exe")
    path.write_bytes(b"x" * (processing.MAX_FILE_SIZE_BYTES + 1))
    with pytest.raises(ValueError, match="5 MB"):
        processing.extract_document(path, path.name)


def test_docx_zip_bomb_rejected_before_opening_xml(tmp_path):
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * (processing.MAX_ARCHIVE_BYTES + 1))
    with pytest.raises(ValueError, match="descomprimir"):
        processing.extract_document(path, path.name)


@pytest.mark.asyncio
async def test_subprocess_runs_offline_and_preserves_unicode(tmp_path, monkeypatch):
    # This also exercises the real 512 MiB / one-thread reader profile.
    monkeypatch.setenv("AWS_SECRET_ACCESS_KEY", "must-not-be-inherited")
    monkeypatch.setenv("OPENAI_API_KEY", "must-not-be-inherited")
    path = tmp_path / "local.txt"
    path.write_text("Atención en Bogotá 😊\n" * 100, encoding="utf-8")
    result = await processing.process_document_locally(
        file_path=str(path),
        filename=path.name,
        retrieval_mode="chunked",
        max_tokens=32,
    )
    assert "".join(chunk["chunk_text"] for chunk in result["chunks"]) == path.read_text(
        encoding="utf-8"
    )


@pytest.mark.asyncio
async def test_ocr_runs_inside_the_real_bounded_child(tmp_path):
    image = Image.new("RGB", (1200, 400), "white")
    draw = ImageDraw.Draw(image)
    draw.text(
        (60, 60),
        "Colombia - pedidos y soporte",
        font=ImageFont.load_default(size=40),
        fill="black",
    )
    draw.text(
        (60, 140),
        "Precio: 25000 pesos",
        font=ImageFont.load_default(size=40),
        fill="black",
    )
    path = tmp_path / "escaneado.pdf"
    image.save(path, "PDF")
    result = await processing.process_document_locally(
        file_path=str(path),
        filename=path.name,
        retrieval_mode="chunked",
        max_tokens=32,
    )
    assert "Colombia" in result["full_text"] and "25000" in result["full_text"]
    assert result["docling_metadata"]["ocr_pages"] == [1]


@pytest.mark.asyncio
async def test_parser_lease_requeues_and_releases_on_failure(monkeypatch):
    redis = AsyncMock()
    redis.set.return_value = False
    reader = AsyncMock(side_effect=ValueError("bad file"))
    monkeypatch.setattr(processing, "process_document_locally", reader)
    with pytest.raises(Retry):
        await processing.process_document_in_queue(redis=redis, file_path="file")
    reader.assert_not_awaited()
    redis.eval.assert_not_awaited()
    redis.set.return_value = True
    with pytest.raises(ValueError, match="bad file"):
        await processing.process_document_in_queue(redis=redis, file_path="file")
    redis.eval.assert_awaited_once()


@pytest.mark.asyncio
async def test_timeout_kills_reader_process_group(monkeypatch):
    process = AsyncMock()
    process.returncode = None
    process.pid = 999999
    process.communicate.side_effect = asyncio.TimeoutError
    monkeypatch.setattr(
        asyncio, "create_subprocess_exec", AsyncMock(return_value=process)
    )
    calls = []
    monkeypatch.setattr(processing.os, "killpg", lambda *args: calls.append(args))
    with pytest.raises(asyncio.TimeoutError):
        await processing.process_document_locally(
            file_path="file",
            filename="file.txt",
            retrieval_mode="full_document",
            max_tokens=32,
        )
    assert calls == [(process.pid, processing.signal.SIGKILL)]
    process.wait.assert_awaited_once()
